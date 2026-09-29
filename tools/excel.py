"""Genera el Excel profesional a partir de data/db.json y permite reimportarlo.

    python tools/excel.py exportar            -> Mi Cinemateca.xlsx
    python tools/excel.py importar [ruta]     -> aplica a db.json lo editado en el Excel
"""
import json
import os
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
from rutas import DB_PATH  # noqa: E402
EST_PATH = ROOT / "data" / "estrenos.json"
XLSX_PATH = ROOT / "Mi Cinemateca.xlsx"

INK = "15151C"
GOLD = "C9962F"
MUTED = "6B6B78"
FONT = "Calibri"

PELI_COLS = [
    # (cabecera, clave, ancho, formato)
    ("ID", "id", 8, None),
    ("Título", "titulo", 40, None),
    ("Título original", "tituloOriginal", 34, None),
    ("Año", "anio", 7, "0"),
    ("Década", None, 8, "0"),
    ("Duración", "duracion", 10, '0" min"'),
    ("Nota", "nota", 7, "0.0"),
    ("Director", "director", 30, None),
    ("País", "pais", 16, None),
    ("Géneros", "generos", 30, None),
    ("Saga", "saga", 22, None),
    ("Fase UCM", "fase", 26, None),
    ("Visto el", "fechaVisto", 12, "dd/mm/yyyy"),
    ("Dónde", "lugar", 26, None),
    ("Favorita", "favorita", 9, None),
    ("Reseña", "resena", 40, None),
    ("Apertura EE.UU. ($)", "apertura", 16, '#,##0'),
    ("Taquilla EE.UU. ($)", "domestica", 16, '#,##0'),
    ("Internacional ($)", "internacional", 16, '#,##0'),
    ("Mundial ($)", None, 16, '#,##0'),
    ("Presupuesto ($)", "presupuesto", 15, '#,##0'),
    ("IMDb", "imdb", 8, None),
    ("FilmAffinity", "filmaffinity", 12, None),
    ("Letterboxd", "letterboxd", 11, None),
]
SERIE_COLS = [("ID", "id", 8, None), ("Título", "titulo", 42, None), ("Años", "anios", 13, None),
              ("Tipo", "tipo", 12, None), ("País", "pais", 16, None), ("Animación", "animacion", 11, None),
              ("Nota", "nota", 7, "0.0"), ("Reseña", "resena", 50, None)]


def links(p):
    ids = p.get("ids") or {}
    return {
        "imdb": f"https://www.imdb.com/title/{ids['imdb']}/" if ids.get("imdb") else None,
        "filmaffinity": f"https://www.filmaffinity.com/es/film{ids['filmaffinity']}.html" if ids.get("filmaffinity") else None,
        "letterboxd": f"https://letterboxd.com/film/{ids['letterboxd']}/" if ids.get("letterboxd") else None,
    }


def style_title(ws, text, sub):
    ws["B2"] = text
    ws["B2"].font = Font(name=FONT, size=24, bold=True, color=INK)
    ws["B3"] = sub
    ws["B3"].font = Font(name=FONT, size=10, color=MUTED)
    ws.sheet_view.showGridLines = False


def add_table(ws, name, cols, rows, start_row=1, style="TableStyleMedium15"):
    for c, (h, _, w, _) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    for r, vals in enumerate(rows, start_row + 1):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            fmt = cols[c - 1][3]
            if fmt:
                cell.number_format = fmt
    last = start_row + max(len(rows), 1)
    ref = f"A{start_row}:{get_column_letter(len(cols))}{last}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    ws.row_dimensions[start_row].height = 22
    return last


def export(db=None, estrenos=None, path=XLSX_PATH):
    db = db or json.loads(DB_PATH.read_text(encoding="utf-8"))
    estrenos = estrenos or (json.loads(EST_PATH.read_text(encoding="utf-8")) if EST_PATH.exists() else {"estrenos": []})
    pelis = sorted(db["peliculas"], key=lambda p: (-(p.get("anio") or 0), -(p.get("nota") or 0)))
    wb = Workbook()

    # ------------------------------------------------------------ Películas
    ws = wb.active
    ws.title = "Películas"
    rows = []
    for i, p in enumerate(pelis, 2):
        t = p.get("taquilla") or {}
        lk = links(p)
        fv = p.get("fechaVisto")
        try:
            fv = datetime.strptime(fv, "%Y-%m-%d").date() if fv else None
        except ValueError:
            pass
        rows.append([
            p["id"], p["titulo"], p.get("tituloOriginal"), p.get("anio"), f'=IF(D{i}="","",FLOOR(D{i},10))',
            p.get("duracion"), p.get("nota"), p.get("director"), p.get("pais"), ", ".join(p.get("generos") or []),
            p.get("saga"), p.get("fase"), fv, p.get("lugar"), "★" if p.get("favorita") else "",
            p.get("resena") or None, t.get("apertura"), t.get("domestica"), t.get("internacional"),
            f'=IF(AND(R{i}="",S{i}=""),"",N(R{i})+N(S{i}))', p.get("presupuesto"),
            "IMDb" if lk["imdb"] else None, "FilmAffinity" if lk["filmaffinity"] else None,
            "Letterboxd" if lk["letterboxd"] else None,
        ])
    last = add_table(ws, "Peliculas", PELI_COLS, rows)
    link_font = Font(name=FONT, color="2F6FB0", underline="single")
    for i, p in enumerate(pelis, 2):
        lk = links(p)
        for col, key in ((22, "imdb"), (23, "filmaffinity"), (24, "letterboxd")):
            if lk[key]:
                c = ws.cell(row=i, column=col)
                c.hyperlink = lk[key]
                c.font = link_font
        ws.cell(row=i, column=7).font = Font(name=FONT, bold=True)
        ws.cell(row=i, column=15).font = Font(name=FONT, color=GOLD)
    ws.freeze_panes = "C2"
    ws.conditional_formatting.add(f"G2:G{last}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B",
        mid_type="num", mid_value=5.5, mid_color="FFE08A",
        end_type="num", end_value=10, end_color="5DBB7A"))
    ws.conditional_formatting.add(f"T2:T{last}", DataBarRule(start_type="min", end_type="max", color="C9962F"))

    # ------------------------------------------------------------ Series
    ws2 = wb.create_sheet("Series")
    srows = [[s["id"], s["titulo"], s.get("anios"), s.get("tipo"), s.get("pais"),
              "Sí" if s.get("animacion") else "", s.get("nota"), s.get("resena") or None]
             for s in sorted(db["series"], key=lambda s: (s.get("tipo") or "", -(s.get("nota") if s.get("nota") is not None else -1)))]
    slast = add_table(ws2, "Series", SERIE_COLS, srows)
    ws2.freeze_panes = "C2"
    ws2.conditional_formatting.add(f"G2:G{slast}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=5.5,
        mid_color="FFE08A", end_type="num", end_value=10, end_color="5DBB7A"))

    # ------------------------------------------------------------ UCM
    ws3 = wb.create_sheet("UCM")
    ucm = sorted([p for p in pelis if p.get("saga") == "UCM"], key=lambda p: (p.get("fase") or "", p.get("anio") or 0))
    ucols = [("Fase", "fase", 30, None), ("Año", "anio", 7, "0"), ("Título", "titulo", 42, None),
             ("Director", "director", 28, None), ("Nota", "nota", 7, "0.0"), ("Presupuesto ($)", None, 16, "#,##0"),
             ("Apertura ($)", None, 15, "#,##0"), ("Mundial ($)", None, 16, "#,##0"),
             ("Beneficio bruto ($)", None, 18, '#,##0;[Red]-#,##0'), ("Rentabilidad", None, 12, '0.0"x"')]
    urows = []
    for i, p in enumerate(ucm, 2):
        t = p.get("taquilla") or {}
        urows.append([p.get("fase"), p.get("anio"), p["titulo"], p.get("director"), p.get("nota"),
                      p.get("presupuesto"), t.get("apertura"), t.get("mundial"),
                      f'=IF(OR(F{i}="",H{i}=""),"",H{i}-F{i})', f'=IF(OR(F{i}="",H{i}=""),"",H{i}/F{i})'])
    ulast = add_table(ws3, "UCM", ucols, urows)
    ws3.conditional_formatting.add(f"E2:E{ulast}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=5.5,
        mid_color="FFE08A", end_type="num", end_value=10, end_color="5DBB7A"))
    ws3.conditional_formatting.add(f"J2:J{ulast}", DataBarRule(start_type="min", end_type="max", color="5DBB7A"))
    # resumen por fase
    fases = sorted({p.get("fase") for p in ucm if p.get("fase")})
    r0 = ulast + 3
    ws3.cell(row=r0, column=1, value="Resumen por fase").font = Font(name=FONT, bold=True, size=13, color=INK)
    hdr = ["Fase", "Películas", "Nota media", "Presupuesto total", "Taquilla mundial", "Rentabilidad"]
    for c, h in enumerate(hdr, 1):
        cell = ws3.cell(row=r0 + 1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
    for k, f in enumerate(fases, r0 + 2):
        ws3.cell(row=k, column=1, value=f)
        ws3.cell(row=k, column=2, value=f'=COUNTIF(UCM[Fase],A{k})')
        ws3.cell(row=k, column=3, value=f'=IFERROR(AVERAGEIF(UCM[Fase],A{k},UCM[Nota]),"")').number_format = "0.0"
        ws3.cell(row=k, column=4, value=f'=SUMIF(UCM[Fase],A{k},UCM[Presupuesto ($)])').number_format = "#,##0"
        ws3.cell(row=k, column=5, value=f'=SUMIF(UCM[Fase],A{k},UCM[Mundial ($)])').number_format = "#,##0"
        ws3.cell(row=k, column=6, value=f'=IFERROR(E{k}/D{k},"")').number_format = '0.0"x"'
    ws3.freeze_panes = "D2"

    # ------------------------------------------------------------ Pendientes
    ws4 = wb.create_sheet("Pendientes")
    pcols = [("Título", "titulo", 44, None), ("Año", "anio", 7, "0"), ("Motivo", "motivo", 40, None),
             ("Añadido", "añadido", 18, None)]
    prow = [[w["titulo"], w.get("anio"), w.get("motivo"), (w.get("añadido") or "")[:10]] for w in db.get("pendientes", [])]
    add_table(ws4, "Pendientes", pcols, prow)

    # ------------------------------------------------------------ Estrenos
    ws5 = wb.create_sheet("Estrenos España")
    ecols = [("Fecha", None, 12, "dd/mm/yyyy"), ("Título", None, 42, None), ("Título original", None, 34, None),
             ("Director", None, 28, None), ("Géneros", None, 30, None), ("País", None, 16, None), ("Destacado", None, 10, None)]
    today = date.today().isoformat()
    erow = [[datetime.strptime(e["fecha"], "%Y-%m-%d").date(), e["titulo"], e.get("original"), e.get("director"),
             ", ".join(e.get("generos") or []), e.get("pais"), "★" if e.get("destacado") else ""]
            for e in estrenos.get("estrenos", []) if e["fecha"] >= today]
    add_table(ws5, "Estrenos", ecols, erow)
    ws5.freeze_panes = "C2"

    # ------------------------------------------------------------ Añadir (como el Excel de siempre)
    wa = wb.create_sheet("Añadir", 0)
    wa.sheet_view.showGridLines = False
    wa["A1"] = "Añade aquí las películas que veas"
    wa["A1"].font = Font(name=FONT, size=18, bold=True, color=INK)
    wa["A2"] = ("Escribe una fila por película (como en tu Excel de siempre), guarda y cierra. "
                "La app las pasa a tu colección, completa director, país, géneros y póster si los dejas vacíos, y vacía esta hoja.")
    wa["A2"].font = Font(name=FONT, size=10, color=MUTED)
    wa["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    wa.merge_cells("A2:J2")
    wa.row_dimensions[2].height = 32
    add_cols = [("Año", None, 8, "0"), ("Duración", None, 10, "0"), ("Nota", None, 8, "0.0"), ("Título", None, 38, None),
                ("Director", None, 26, None), ("País", None, 16, None), ("Géneros", None, 24, None),
                ("Dónde", None, 24, None), ("Visto el", None, 12, "dd/mm/yyyy"), ("Reseña", None, 40, None)]
    add_table(wa, "Anadir", add_cols, [[None] * len(add_cols) for _ in range(25)], start_row=4, style="TableStyleMedium2")
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="10", allow_blank=True,
                        showErrorMessage=True, errorTitle="Nota", error="La nota va de 0 a 10")
    wa.add_data_validation(dv)
    dv.add("C5:C29")
    wa.freeze_panes = "A5"

    # ------------------------------------------------------------ Resumen
    wr = wb.create_sheet("Resumen", 1)
    style_title(wr, "Mi Cinemateca", f"Actualizado el {datetime.now().strftime('%d/%m/%Y %H:%M')} · generado desde la app")
    wr.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFGHIJKLMN", (26, 12, 12, 3, 30, 12, 12, 3, 26, 12, 12, 3, 12)):
        wr.column_dimensions[col].width = w
    kpis = [
        ("Películas vistas", "=COUNTA(Peliculas[Título])", "0"),
        ("Nota media", "=AVERAGE(Peliculas[Nota])", "0.00"),
        ("Horas de cine", "=SUM(Peliculas[Duración])/60", "#,##0"),
        ("Obras maestras (≥9)", '=COUNTIF(Peliculas[Nota],">=9")', "0"),
        ("Series vistas", "=COUNTA(Series[Título])", "0"),
        ("Nota media series", "=AVERAGE(Series[Nota])", "0.00"),
    ]
    kfill = PatternFill("solid", fgColor=INK)
    for n, (lab, f, fmt) in enumerate(kpis):
        col = 2 + (n % 3) * 4
        row = 5 + (n // 3) * 3
        a = wr.cell(row=row, column=col, value=lab)
        a.font = Font(name=FONT, size=9, color="BDBDC7")
        a.fill = kfill
        wr.cell(row=row, column=col + 1).fill = kfill
        wr.cell(row=row, column=col + 2).fill = kfill
        b = wr.cell(row=row + 1, column=col, value=f)
        b.font = Font(name=FONT, size=20, bold=True, color=GOLD)
        b.number_format = fmt
        b.alignment = Alignment(horizontal="left")
        b.fill = kfill
        wr.cell(row=row + 1, column=col + 1).fill = kfill
        wr.cell(row=row + 1, column=col + 2).fill = kfill

    def block(r, c, title, headers, data, fmts):
        wr.cell(row=r, column=c, value=title).font = Font(name=FONT, bold=True, size=12, color=INK)
        for j, h in enumerate(headers):
            cell = wr.cell(row=r + 1, column=c + j, value=h)
            cell.font = Font(name=FONT, bold=True, size=9, color=MUTED)
            cell.border = Border(bottom=Side(style="thin", color="CCCCCC"))
        for i, rowv in enumerate(data, r + 2):
            for j, v in enumerate(rowv):
                cell = wr.cell(row=i, column=c + j, value=v)
                cell.font = Font(name=FONT, size=10)
                if fmts[j]:
                    cell.number_format = fmts[j]
        return r + 2 + len(data)

    # por género (fórmulas vivas)
    gcount = Counter(g for p in pelis for g in (p.get("generos") or []))
    R = 12
    gdata = []
    for k, (g, _) in enumerate(gcount.most_common(), R + 2):
        gdata.append([g, f'=COUNTIF(Peliculas[Géneros],"*"&B{k}&"*")',
                      f'=IFERROR(AVERAGEIF(Peliculas[Géneros],"*"&B{k}&"*",Peliculas[Nota]),"")'])
    gend = block(R, 2, "Por género", ["Género", "Películas", "Nota media"], gdata, [None, "0", "0.0"])
    wr.conditional_formatting.add(f"D{R+2}:D{gend-1}", ColorScaleRule(
        start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFE08A",
        end_type="max", end_color="5DBB7A"))

    # por década
    decs = sorted({(p["anio"] // 10) * 10 for p in pelis if p.get("anio")}, reverse=True)
    ddata = []
    for k, d in enumerate(decs, R + 2):
        ddata.append([d, f"=COUNTIF(Peliculas[Década],F{k})", f'=IFERROR(AVERAGEIF(Peliculas[Década],F{k},Peliculas[Nota]),"")'])
    dend = block(R, 6, "Por década", ["Década", "Películas", "Nota media"], ddata, ["0\"s\"", "0", "0.0"])
    wr.conditional_formatting.add(f"H{R+2}:H{dend-1}", ColorScaleRule(
        start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFE08A",
        end_type="max", end_color="5DBB7A"))

    # por país
    pc = Counter(p.get("pais") for p in pelis if p.get("pais"))
    pdata = []
    for k, (pa, _) in enumerate(pc.most_common(12), R + 2):
        pdata.append([pa, f"=COUNTIF(Peliculas[País],J{k})", f'=IFERROR(AVERAGEIF(Peliculas[País],J{k},Peliculas[Nota]),"")'])
    block(R, 10, "Por país (top 12)", ["País", "Películas", "Nota media"], pdata, [None, "0", "0.0"])

    # directores
    dirs = defaultdict(list)
    for p in pelis:
        for d in (p.get("director") or "").split(" / "):
            if d and p.get("nota") is not None:
                dirs[d].append(p["nota"])
    top_dirs = sorted([(d, len(v), sum(v) / len(v)) for d, v in dirs.items() if len(v) >= 3], key=lambda x: -x[2])
    R2 = max(gend, dend) + 2
    ddata2 = [[d, n, round(a, 2)] for d, n, a in top_dirs[:20]]
    dirend = block(R2, 2, "Tus directores (mín. 3 películas)", ["Director", "Películas", "Nota media"], ddata2, [None, "0", "0.00"])

    # top 20
    top = sorted([p for p in pelis if p.get("nota") is not None], key=lambda p: (-p["nota"], p.get("anio") or 0))[:20]
    block(R2, 6, "Tu top 20", ["Título", "Año", "Nota"], [[p["titulo"], p.get("anio"), p["nota"]] for p in top], [None, "0", "0.0"])
    wr.column_dimensions["F"].width = 34

    # distribución de notas
    bins = [[f"{b}–{b + 1}", f'=COUNTIFS(Peliculas[Nota],">={b}",Peliculas[Nota],"<{b + 1 if b < 9 else 10.01}")'] for b in range(10)]
    wr.cell(row=R2, column=10, value="Distribución de notas").font = Font(name=FONT, bold=True, size=12, color=INK)
    for j, h in enumerate(["Rango", "Películas"]):
        wr.cell(row=R2 + 1, column=10 + j, value=h).font = Font(name=FONT, bold=True, size=9, color=MUTED)
    for i, (a, b) in enumerate(bins, R2 + 2):
        wr.cell(row=i, column=10, value=a)
        wr.cell(row=i, column=11, value=b)

    # gráficos nativos
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = "Películas por género"
    ch.y_axis.title = None
    ch.add_data(Reference(wr, min_col=3, min_row=R + 1, max_row=gend - 1), titles_from_data=True)
    ch.set_categories(Reference(wr, min_col=2, min_row=R + 2, max_row=gend - 1))
    ch.legend = None
    ch.height, ch.width = 9, 14
    ch.series[0].graphicalProperties.solidFill = GOLD
    wr.add_chart(ch, f"N{R}")
    ch2 = BarChart()
    ch2.style = 10
    ch2.title = "Distribución de tus notas"
    ch2.add_data(Reference(wr, min_col=11, min_row=R2 + 1, max_row=R2 + 11), titles_from_data=True)
    ch2.set_categories(Reference(wr, min_col=10, min_row=R2 + 2, max_row=R2 + 11))
    ch2.legend = None
    ch2.height, ch2.width = 7.5, 14
    ch2.series[0].graphicalProperties.solidFill = INK
    wr.add_chart(ch2, f"N{R + 20}")

    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0

    # guardado atómico (si el Excel está abierto, avisa en vez de romperlo)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(Path(path).parent))
    os.close(fd)
    wb.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError:
        os.remove(tmp)
        raise PermissionError("El Excel está abierto: ciérralo para que se pueda actualizar.")
    return str(path)


# ---------------------------------------------------------------- importar
def _val(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def import_pro(path=XLSX_PATH, db=None):
    """Aplica al db lo editado a mano en el Excel profesional. Devuelve (db, resumen)."""
    db = db or json.loads(DB_PATH.read_text(encoding="utf-8"))
    wb = load_workbook(path, data_only=True)
    ws = wb["Películas"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    by_id = {p["id"]: p for p in db["peliculas"]}
    changed = added = 0
    import random
    marca = datetime.utcnow().isoformat(timespec="seconds")
    for row in ws.iter_rows(min_row=2, values_only=True):
        g = lambda h: _val(row[idx[h]]) if h in idx and idx[h] < len(row) else None
        if not g("Título"):
            continue
        fv = g("Visto el")
        if isinstance(fv, datetime):
            fv = fv.date().isoformat()
        upd = {
            "titulo": g("Título"), "tituloOriginal": g("Título original"),
            "anio": int(g("Año")) if g("Año") else None,
            "duracion": int(g("Duración")) if g("Duración") else None,
            "nota": round(float(g("Nota")), 1) if g("Nota") is not None else None,
            "director": g("Director") or "", "pais": g("País") or "",
            "generos": [x.strip() for x in (g("Géneros") or "").split(",") if x.strip()],
            "saga": g("Saga"), "fase": g("Fase UCM"), "fechaVisto": fv, "lugar": g("Dónde"),
            "favorita": bool(g("Favorita")), "resena": g("Reseña") or "",
            "presupuesto": g("Presupuesto ($)"),
        }
        taq = {"apertura": g("Apertura EE.UU. ($)"), "domestica": g("Taquilla EE.UU. ($)"),
               "internacional": g("Internacional ($)")}
        pid = g("ID")
        if pid and pid in by_id:
            p = by_id[pid]
            before = json.dumps(p, sort_keys=True, ensure_ascii=False)
            p.update(upd)
            t = p.setdefault("taquilla", {})
            t.update(taq)
            t["mundial"] = (t.get("domestica") or 0) + (t.get("internacional") or 0) if (t.get("domestica") or t.get("internacional")) else None
            if json.dumps(p, sort_keys=True, ensure_ascii=False) != before:
                p["mod"] = marca
                changed += 1
        else:
            p = {**upd, "id": f"p{datetime.utcnow():%Y%m%d%H%M%S}{random.randint(0, 999999):06d}",
                 "taquilla": {**taq, "mundial": None}, "ids": {}, "poster": None, "origen": "excel-nueva",
                 "mod": marca, "añadido": marca}
            db["peliculas"].append(p)
            added += 1
    # hoja «Añadir»: filas nuevas con las columnas del Excel de siempre
    if "Añadir" in wb.sheetnames:
        wa = wb["Añadir"]
        hdr_row = next((r for r in range(1, 10) if wa.cell(row=r, column=4).value == "Título"), None)
        if hdr_row:
            hdr = [c.value for c in wa[hdr_row]]
            ix = {h: i for i, h in enumerate(hdr) if h}
            existentes = {(_norm(p["titulo"]), p.get("anio")) for p in db["peliculas"]}
            existentes |= {(_norm(p.get("tituloOriginal")), p.get("anio")) for p in db["peliculas"] if p.get("tituloOriginal")}
            for row in wa.iter_rows(min_row=hdr_row + 1, values_only=True):
                g = lambda h: _val(row[ix[h]]) if h in ix and ix[h] < len(row) else None
                titulo = g("Título")
                if not titulo:
                    continue
                titulo = str(titulo).strip()
                if titulo.isupper():
                    titulo = _title_case(titulo)
                anio = int(g("Año")) if g("Año") else None
                if (_norm(titulo), anio) in existentes:
                    continue
                fv = g("Visto el")
                if isinstance(fv, datetime):
                    fv = fv.date().isoformat()
                nota = g("Nota")
                if isinstance(nota, str):
                    nota = float(nota.replace(",", "."))
                director = g("Director") or ""
                db["peliculas"].append({
                    "id": f"p{datetime.utcnow():%Y%m%d%H%M%S}{random.randint(0, 999999):06d}",
                    "titulo": titulo, "tituloOriginal": None, "anio": anio,
                    "duracion": int(g("Duración")) if g("Duración") else None,
                    "nota": round(float(nota), 1) if nota is not None else None,
                    "director": _title_case(director) if str(director).isupper() else director,
                    "pais": _title_case(str(g("País"))).replace(" De ", " de ").replace(" Del ", " del ") if g("País") else "",
                    "generos": [x.strip().capitalize() for x in str(g("Géneros") or "").split(",") if x.strip()],
                    "saga": None, "fase": None, "presupuesto": None, "taquilla": {},
                    "fechaVisto": fv, "lugar": g("Dónde"), "resena": g("Reseña") or "", "favorita": False,
                    "ids": {}, "poster": None, "origen": "excel-nueva", "mod": marca, "añadido": marca,
                })
                existentes.add((_norm(titulo), anio))
                added += 1
    return db, {"modificadas": changed, "nuevas": added}


def _norm(x):
    import re
    import unicodedata
    x = unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", x.upper())


def _title_case(x):
    try:
        from importar_excel import title_case
        return title_case(x)
    except Exception:
        return str(x).title()


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "exportar"
    if cmd == "exportar":
        print("Excel generado:", export())
    elif cmd == "importar":
        db, res = import_pro(sys.argv[2] if len(sys.argv) > 2 else XLSX_PATH)
        DB_PATH.write_text(json.dumps(db, ensure_ascii=False, indent=1), encoding="utf-8")
        print("Importado:", res)
