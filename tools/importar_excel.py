"""Importa el Excel original (Películas.xlsx) a data/db.json.

Uso:  python tools/importar_excel.py "C:/ruta/Películas.xlsx"

Aplica correcciones de datos, normaliza mayúsculas, añade géneros y títulos
originales (tools/generos_map.txt), detecta sagas y fusiona la hoja UCM.
"""
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from rutas import DB_PATH  # noqa: E402

GENEROS = {
    "A": "Acción", "Av": "Aventura", "An": "Animación", "CF": "Ciencia ficción",
    "C": "Comedia", "Cr": "Crimen", "D": "Drama", "F": "Fantasía", "T": "Terror",
    "Th": "Thriller", "R": "Romance", "M": "Musical", "B": "Bélico", "W": "Western",
    "S": "Superhéroes", "Mi": "Misterio", "Bio": "Biográfico", "H": "Histórico",
    "Fa": "Familiar", "Doc": "Documental",
}


def norm(s):
    """Clave de comparación: sin acentos, sin signos, mayúsculas."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", s.upper())


# --- Correcciones detectadas en el Excel original -------------------------
TITULO_FIX = {
    "UNTILDOWN": "UNTIL DAWN",
    "SACRYMOVIE5": "SCARY MOVIE 5",
    "PROYECTX": "PROJECT X",
    "QUIENESTAMATANDOALOSMONECOS": "¿QUIÉN ESTÁ MATANDO A LOS MUÑECOS?",
    "LOSPINGUINOSDELSRPOPER": "LOS PINGÜINOS DEL SR. POPPER",
    "JUMANJISIGUENTENIVEL": "JUMANJI: SIGUIENTE NIVEL",
    "ICEAGE3ELORIGENDELOSDINOSAURIOS": "ICE AGE 3: EL ORIGEN DE LOS DINOSAURIOS",
}
DIRECTOR_FIX_POR_TITULO = {
    "THORRAGNAROK": "TAIKA WAITITI",
    "GUARDIANESDELAGALAXIAVOL2": "JAMES GUNN",
    "ELBUENPATRON": "FERNANDO LEÓN DE ARANOA",
    "THEINNOCENTS": "ESKIL VOGT",
    "SPIDERMANNOWAYHOME": "JON WATTS",
    "UNTILDAWN": "DAVID F. SANDBERG",
    "MATRIX": "LANA WACHOWSKI / LILLY WACHOWSKI",
}
PAIS_FIX_POR_TITULO = {
    "ELBUENPATRON": "ESPAÑA",
    "THEINNOCENTS": "NORUEGA",
    "SPIDERMANNOWAYHOME": "ESTADOS UNIDOS",
    "LABELLAYLABESTIA": "ESTADOS UNIDOS",
    "ZOMBIESPARTY": "REINO UNIDO",
    "NOTTINGHILL": "REINO UNIDO",
}
DIRECTOR_TYPO = {
    "JAKE SCHREIRER": "JAKE SCHREIER", "KJAKE KASDAN": "JAKE KASDAN",
    "JON TURTLELTAUB": "JON TURTELTAUB", "M. NICHT SHYMALAN": "M. NIGHT SHYAMALAN",
    "JONATHAM DEMME": "JONATHAN DEMME", "MCK JACKSON": "MICK JACKSON",
    "STEPHEN SOMMERA": "STEPHEN SOMMERS", "JOHN SIGLETON": "JOHN SINGLETON",
    "ROGER MITCHELL": "ROGER MICHELL", "GENNEDY TARTAKOVSKY": "GENNDY TARTAKOVSKY",
    "VICENTE VILANUEVA": "VICENTE VILLANUEVA", "DANIEL CALPALSORO": "DANIEL CALPARSORO",
    "ALEJANDRO AMENÁNABAR": "ALEJANDRO AMENÁBAR", "HELENNE ZIMMER": "HÉLÈNE ZIMMER",
    "CHRIS WILIAMS": "CHRIS WILLIAMS", "MATT REAVES": "MATT REEVES",
    "JAVIER RUÍZ CALDERA": "JAVIER RUIZ CALDERA", "LLUÍS QUILEZ": "LLUÍS QUÍLEZ",
    "GARETH JENNINGS": "GARTH JENNINGS", "RUBEN OSTLUND": "RUBEN ÖSTLUND",
    "DAVID F. SNADBERG": "DAVID F. SANDBERG", "HERMANAS WACHOWSKI": "LANA WACHOWSKI / LILLY WACHOWSKI",
    "JOHN FRANCIS DALEY / JONATHAN GOLDSTEIN": "JOHN FRANCIS DALEY / JONATHAN GOLDSTEIN",
}
PAIS_TYPO = {"EPAÑA": "ESPAÑA"}

SAGAS = [
    (r"^STAR WARS|^HAN SOLO|^ROGUE ONE", "Star Wars"),
    (r"SPIDER-MAN", "Spider-Man"),
    (r"^EXPEDIENTE WARREN|^LA MONJA", "Universo Expediente Warren"),
    (r"^INSIDIOUS", "Insidious"),
    (r"^LA (PRIMERA )?PURGA", "La Purga"),
    (r"^SAW\b", "Saw"),
    (r"^SCARY MOVIE", "Scary Movie"),
    (r"^REC\b", "REC"),
    (r"^TORRENTE", "Torrente"),
    (r"SHIN CHAN", "Shin Chan"),
    (r"^DORAEMON", "Doraemon"),
    (r"^POKÉMON", "Pokémon"),
    (r"^DETECTIVE CONAN", "Detective Conan"),
    (r"^DRAGON BALL", "Dragon Ball"),
    (r"^TOY STORY", "Toy Story"),
    (r"^ICE AGE", "Ice Age"),
    (r"^FAST & FURIOUS", "Fast & Furious"),
    (r"^EL SEÑOR DE LOS ANILLOS", "El Señor de los Anillos"),
    (r"^CREPÚSCULO", "Crepúsculo"),
    (r"RESACÓN|^R3SACÓN", "Resacón"),
    (r"^GRU\b", "Gru"),
    (r"^MADAGASCAR", "Madagascar"),
    (r"^JUMANJI", "Jumanji"),
    (r"^KINGSMAN", "Kingsman"),
    (r"^ROCKY|^CREED", "Rocky / Creed"),
    (r"^ALIEN", "Alien"),
    (r"^MEN IN BLACK", "Men in Black"),
    (r"^AMERICAN PIE", "American Pie"),
    (r"^STRANGERS", "The Strangers"),
    (r"^FELIZ DÍA DE TU MUERTE", "Feliz día de tu muerte"),
    (r"^IT\b|^IT CAPÍTULO", "It"),
    (r"^EL PADRINO", "El Padrino"),
    (r"^BATMAN BEGINS|^EL CABALLERO OSCURO|^THE BATMAN", "Batman"),
    (r"^PUÑALES POR LA ESPALDA", "Puñales por la espalda"),
    (r"^INSIDE OUT", "Inside Out"),
    (r"^SUPER MARIO", "Super Mario"),
    (r"^CARS\b", "Cars"),
    (r"^OCHO APELLIDOS", "Ocho apellidos"),
    (r"^DOS TONTOS", "Dos tontos muy tontos"),
    (r"^NIÑOS GRANDES", "Niños grandes"),
    (r"^BITELCHÚS", "Bitelchús"),
    (r"^JURASSIC", "Jurassic Park"),
    (r"^DUNE", "Dune"),
    (r"^BLADE RUNNER", "Blade Runner"),
    (r"^LOS 4 FANTÁSTICOS$|^LOS 4 FANTÁSTICOS Y SILVER", "Los 4 Fantásticos (2005)"),
    (r"^MISIÓN IMPOSIBLE", "Misión Imposible"),
    (r"^LA MOMIA", "La Momia"),
]

SMALL = {"de", "del", "la", "las", "el", "los", "y", "e", "en", "a", "al", "por", "con",
         "un", "una", "o", "u", "para", "sin", "sus", "su", "the", "of", "and", "on", "in",
         "at", "to", "for", "vs", "vs."}
KEEP_UPPER = {"II", "III", "IV", "V", "VI", "UCM", "F1", "REC", "X", "G.I.", "E.T.",
              "S.A.", "KKKLAN", "NY", "USA", "UK", "XY9", "UFO", "JR", "JR."}
NAME_PARTICLES = {"de", "del", "la", "van", "von", "der", "dos", "da", "di", "le", "y"}
SPECIAL_WORDS = {"JOON-HO": "Joon-ho", "CHAN-WOOK": "Chan-wook", "MCQUARRIE": "McQuarrie",
                 "MACDONALD": "MacDonald", "MACFARLANE": "MacFarlane", "R3SACÓN": "R3sacón",
                 "KKKLAN": "KKKlan", "IÑÁRRITU": "Iñárritu", "D'": "d'",
                 "WALL-E": "WALL·E", "DACOSTA": "DaCosta", "ØVREDAL": "Øvredal"}


def cap_word(w, particles=None, first=False):
    up = w.upper()
    if up in SPECIAL_WORDS:
        return SPECIAL_WORDS[up]
    core = up.rstrip(":.,;!?")
    if core in KEEP_UPPER:
        return core + up[len(core):]
    low = w.lower()
    if not first and particles and low in particles:
        return low
    if re.fullmatch(r"[\d.,]+", w):
        return w
    if up.startswith("MC") and len(up) > 3:
        return "Mc" + up[2].upper() + up[3:].lower()
    # capitaliza tras signos iniciales (¿ ¡ ( ") y tras guion/punto
    out, cap_next = [], True
    for ch in low:
        if cap_next and ch.isalpha():
            out.append(ch.upper())
            cap_next = False
        else:
            out.append(ch)
        if ch in "-.":
            cap_next = True
    return "".join(out)


def title_case(s, particles=SMALL):
    s = re.sub(r"\s+", " ", str(s).strip())
    words = s.split(" ")
    res = []
    for i, w in enumerate(words):
        prev = words[i - 1] if i else ""
        first = i == 0 or prev.endswith((":", ".", "-", "¿", "¡")) or prev in ("-", "–")
        res.append(cap_word(w, particles, first))
    return " ".join(res)


def director_case(s):
    s = re.sub(r"\s*[/,]\s*", " / ", str(s).strip())
    s = re.sub(r"\s+", " ", s)
    out = " / ".join(title_case(p, NAME_PARTICLES) for p in s.split(" / "))
    return out.replace(" de Palma", " De Palma").replace("Joaquim dos Santos", "Joaquim Dos Santos")


def money(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v)
    v = re.sub(r"[^\d]", "", str(v))
    return int(v) if v else None


def load_map():
    m = {}
    for line in (ROOT / "tools" / "generos_map.txt").read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.startswith("#"):
            continue
        key, orig, codes = [p.strip() for p in line.split("|")]
        year = None
        if "@" in key:
            key, year = key.split("@")
            year = int(year)
        m[(norm(key), year)] = (orig, [GENEROS[c] for c in codes.split(",") if c])
    return m


def lookup(m, titulo, anio):
    return m.get((norm(titulo), anio)) or m.get((norm(titulo), None))


def detect_saga(t):
    for rx, name in SAGAS:
        if re.search(rx, t):
            return name
    return None


def read_ucm(ws):
    """Devuelve {norm(titulo): {fase, presupuesto, taquilla}} y los títulos futuros."""
    out, futuros, fase = {}, [], None
    for r in ws.iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * 15
        if isinstance(r[1], str) and r[1].upper().startswith("FASE"):
            fase = re.sub(r"\s+", " ", r[1]).strip().replace(" - ", " · ").capitalize().replace("saga del multiverso", "Saga del Multiverso").replace("saga del infinito", "Saga del Infinito")
            continue
        titulo = r[4]
        if not titulo:
            continue
        info = {
            "fase": fase,
            "presupuesto": money(r[8]),
            "taquilla": {"apertura": money(r[9]), "domestica": money(r[10]), "internacional": money(r[11])},
        }
        k = norm(titulo).replace("LOS4FANTASTICOS", "LOSCUATROFANTASTICOS")
        out[k] = info
        if r[3] is None:
            futuros.append({"titulo": title_case(titulo).replace(": 5", " 5"), "anio": int(r[1]) if isinstance(r[1], (int, float)) else None, "fase": fase})
    return out, futuros


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    gmap = load_map()
    ucm, ucm_futuros = read_ucm(wb["UCM"])

    ws = wb["PELÍCULAS"]
    peliculas, vistos, sin_mapa, correcciones = [], set(), [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        _, anio, dur, nota, titulo, director, pais, estreno, dom, intl, *_ = list(row) + [None] * 12
        if titulo is None:
            continue
        if isinstance(titulo, float):
            titulo = str(int(titulo))
        titulo = re.sub(r"\s+", " ", str(titulo)).strip()
        k = norm(titulo)
        if k in TITULO_FIX:
            correcciones.append(f"Título: {titulo} → {TITULO_FIX[k]}")
            titulo = TITULO_FIX[k]
            k = norm(titulo)
        anio = int(anio) if anio else None
        dedup = (k, anio)
        if dedup in vistos:
            correcciones.append(f"Duplicado eliminado: {titulo} ({anio})")
            continue
        vistos.add(dedup)

        director = re.sub(r"\s+", " ", str(director or "")).strip()
        director = DIRECTOR_TYPO.get(director, director)
        if k in DIRECTOR_FIX_POR_TITULO and director != DIRECTOR_FIX_POR_TITULO[k]:
            correcciones.append(f"Director de {titulo}: {director} → {DIRECTOR_FIX_POR_TITULO[k]}")
            director = DIRECTOR_FIX_POR_TITULO[k]
        pais = re.sub(r"\s+", " ", str(pais or "")).strip()
        pais = PAIS_TYPO.get(pais, pais)
        if k in PAIS_FIX_POR_TITULO and pais != PAIS_FIX_POR_TITULO[k]:
            correcciones.append(f"País de {titulo}: {pais} → {PAIS_FIX_POR_TITULO[k]}")
            pais = PAIS_FIX_POR_TITULO[k]

        found = lookup(gmap, titulo, anio)
        if not found:
            sin_mapa.append(f"{titulo} ({anio})")
        orig, generos = found or (None, [])

        taq = {"apertura": money(estreno), "domestica": money(dom), "internacional": money(intl)}
        saga, fase, presupuesto = detect_saga(titulo), None, None
        u = ucm.get(k)
        if u:
            saga, fase, presupuesto = "UCM", u["fase"], u["presupuesto"]
            for kk, vv in u["taquilla"].items():
                if taq[kk] is None and vv:
                    taq[kk] = vv
        if taq["domestica"] is not None or taq["internacional"] is not None:
            taq["mundial"] = (taq["domestica"] or 0) + (taq["internacional"] or 0)
        else:
            taq["mundial"] = None

        peliculas.append({
            "titulo": orig if orig and norm(orig) == norm(titulo) and pais in ("ESPAÑA", "ARGENTINA") else title_case(titulo),
            "tituloOriginal": orig,
            "anio": anio,
            "duracion": int(dur) if dur else None,
            "nota": round(float(nota), 1) if nota is not None else None,
            "director": director_case(director) if director else "",
            "pais": title_case(pais.lower(), SMALL) if pais else "",
            "generos": generos,
            "saga": saga,
            "fase": fase,
            "presupuesto": presupuesto,
            "taquilla": taq,
            "fechaVisto": None,
            "lugar": None,
            "resena": "",
            "favorita": False,
            "ids": {},
            "poster": None,
            "origen": "excel",
        })

    peliculas.sort(key=lambda p: (-(p["anio"] or 0), -(p["nota"] or 0)))
    for i, p in enumerate(peliculas, 1):
        p["id"] = f"p{i:04d}"

    # --- Series -------------------------------------------------------------
    SERIE_FIX = {
        "ELCOLAPSO": {"pais": "FRANCIA"}, "ANTIDISTURBIOS": {"pais": "ESPAÑA"},
        "LOVEDEATHROBOTS": {"pais": "ESTADOS UNIDOS", "animacion": True},
        "CHERNOBIL": {"animacion": False}, "SARGENTOKERORO": {"pais": "JAPÓN"},
    }
    SERIE_TITULO = {"El hundiminento de Japón": "El hundimiento de Japón",
                    "Cómo se covirtieron en tiranos": "Cómo se convirtieron en tiranos",
                    "Farenheit 9/11": "Fahrenheit 9/11", "Los Picapierdra": "Los Picapiedra",
                    "Star Wars: The Clon Wars": "Star Wars: The Clone Wars",
                    "Putín: De espía a presidente": "Putin: De espía a presidente"}
    series, seen = [], set()
    for r in wb["SERIES "].iter_rows(values_only=True):
        titulo, anios, nota, tipo, pais, anim, *_ = list(r) + [None] * 7
        if not titulo:
            continue
        titulo = SERIE_TITULO.get(str(titulo).strip(), str(titulo).strip())
        k = norm(titulo)
        if k in seen:
            continue
        seen.add(k)
        if isinstance(anios, float):
            anios = str(int(anios))
        anios = re.sub(r"\s*-\s*", "–", str(anios)).replace("act", "act.") if anios else ""
        fx = SERIE_FIX.get(k, {})
        pais = fx.get("pais", pais)
        series.append({
            "id": f"s{len(series) + 1:04d}",
            "titulo": titulo,
            "anios": anios,
            "nota": round(float(nota), 1) if nota is not None else None,
            "tipo": (tipo or "SERIE").capitalize() if tipo else "Serie",
            "pais": title_case(str(pais).lower()) if pais else "",
            "animacion": fx.get("animacion", anim == "ANIMACIÓN"),
            "resena": "",
            "poster": None,
            "ids": {},
        })

    pendientes = []
    for f in ucm_futuros:
        pendientes.append({"id": f"w{len(pendientes) + 1:04d}", "titulo": f["titulo"], "anio": f["anio"],
                           "motivo": f"UCM · {f['fase']}", "añadido": datetime.now().isoformat(timespec="seconds")})

    db = {
        "version": 1,
        "actualizado": datetime.now().isoformat(timespec="seconds"),
        "peliculas": peliculas,
        "series": series,
        "pendientes": pendientes,
        "correcciones": correcciones + [
            "Series: El Colapso → Francia, Antidisturbios → España, Love, Death + Robots → EE.UU. (animación), Sargento Keroro → Japón",
            "Series: erratas en El hundimiento de Japón, Fahrenheit 9/11, Los Picapiedra, Star Wars: The Clone Wars; duplicados de Mirmo y La hora chanante",
            "Directores con erratas corregidos (Shyamalan, Demme, Turteltaub, Sommers, Singleton, Michell, Reeves, Amenábar, Kasdan…)",
        ],
    }
    old = json.loads(DB_PATH.read_text(encoding="utf-8")) if DB_PATH.exists() else {}
    # conserva enriquecimientos previos (ids/póster) si se reimporta
    prev = {(norm(p["titulo"]), p.get("anio")): p for p in old.get("peliculas", [])}
    for p in peliculas:
        o = prev.get((norm(p["titulo"]), p["anio"]))
        if o:
            p["ids"], p["poster"] = o.get("ids", {}), o.get("poster")
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    DB_PATH.write_text(json.dumps(db, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"Películas: {len(peliculas)} · Series: {len(series)} · Pendientes: {len(pendientes)}")
    print(f"Correcciones aplicadas ({len(correcciones)}):")
    for c in correcciones:
        print("  -", c)
    if sin_mapa:
        print(f"SIN GÉNERO ({len(sin_mapa)}):", ", ".join(sin_mapa))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(Path.home() / "Downloads" / "Películas.xlsx"))
